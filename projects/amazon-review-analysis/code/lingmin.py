from sklearn.feature_extraction.text import CountVectorizer
from sklearn.naive_bayes import MultinomialNB
from imblearn.over_sampling import RandomOverSampler
from pathlib import Path

import pandas as pd

ROOT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT_DIR / "data"
OUTPUT_DIR = ROOT_DIR / "outputs"
OUTPUT_DIR.mkdir(exist_ok=True)

data = pd.read_csv(DATA_DIR / "app4gai3.csv")
data = data.dropna(subset=['reviewText'])
data = data.dropna(subset=['summary'])

comments = data['reviewText'].tolist()
ratings = data['overall'].tolist()
summary = data['summary'].tolist()

vectorizer1 = CountVectorizer()
vectorizer2 = CountVectorizer()
X1 = vectorizer1.fit_transform(comments)
X2 = vectorizer2.fit_transform(summary)

ros = RandomOverSampler(random_state=0)
X1_resampled, y1_resampled = ros.fit_resample(X1, ratings)
X2_resampled, y2_resampled = ros.fit_resample(X2, ratings)

clf1 = MultinomialNB()
clf1.fit(X1_resampled, y1_resampled)
clf2 = MultinomialNB()
clf2.fit(X2_resampled, y2_resampled)

test_data = pd.read_csv(DATA_DIR / "app6.csv")
test_data = test_data.dropna(subset=['reviewText'])
test_data = test_data.dropna(subset=['summary'])
test_comments = test_data['reviewText'].tolist()
test_summary = test_data['summary'].tolist()

X_test_c = vectorizer1.transform(test_comments)
X_test_s = vectorizer2.transform(test_summary)

predicted_ratings1 = clf1.predict(X_test_c)
predicted_ratings2 = clf2.predict(X_test_s)

predicted_ratings = []
for i in range(0,len(predicted_ratings1)):
    predicted_ratings.append(predicted_ratings1[i]*0.5+predicted_ratings2[i]*0.5)
from openpyxl import load_workbook
wb = load_workbook(DATA_DIR / "app6yu4.xlsx")
ws = wb.active
for row, entry in enumerate(predicted_ratings, start=2):
    ws.cell(row=row, column=4, value=entry)
wb.save(OUTPUT_DIR / "app6yu4.xlsx")
sum = 0
c=test_data['overall'].tolist()
for i in range(0, len(predicted_ratings)):
    sum = sum + abs(predicted_ratings[i] - c[i])
inaccuracy = sum / len(predicted_ratings)
print(inaccuracy)
