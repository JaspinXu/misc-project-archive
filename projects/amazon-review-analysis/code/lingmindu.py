from sklearn.feature_extraction.text import CountVectorizer
from sklearn.naive_bayes import MultinomialNB
from sklearn.model_selection import GridSearchCV
from pathlib import Path

import pandas as pd
import random

ROOT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT_DIR / "data"
OUTPUT_DIR = ROOT_DIR / "outputs"
OUTPUT_DIR.mkdir(exist_ok=True)

train_data = pd.read_csv(DATA_DIR / "app4444.csv")
train_data = train_data.dropna(subset=['reviewText'])
train_data = train_data.reset_index(drop=True)
train_data = train_data.dropna(subset=['summary'])
train_data = train_data.reset_index(drop=True)
train_comments = []
train_ratings = []
train_summary = []
for i in range(0, len(train_data)):
    train_comments.append(train_data['reviewText'][i])
    train_ratings.append(train_data['overall'][i])
    train_summary.append(train_data['summary'][i])
test_comments = []
test_ratings = []
test_summary = []
test_data = pd.read_csv(DATA_DIR / "app6.csv")
test_data = test_data.dropna(subset=['reviewText'])
test_data = test_data.reset_index(drop=True)
test_data = test_data.dropna(subset=['summary'])
test_data = test_data.reset_index(drop=True)
for i in range(0, len(test_data)):
    test_comments.append(test_data['reviewText'][i])
    test_ratings.append(test_data['overall'][i])
    test_summary.append(test_data['summary'][i])
vectorizer_comments = CountVectorizer()
X_train_comments = vectorizer_comments.fit_transform(train_comments)
X_test_comments = vectorizer_comments.transform(test_comments)

vectorizer_summary = CountVectorizer()
X_train_summary = vectorizer_summary.fit_transform(train_summary)
X_test_summary = vectorizer_summary.transform(test_summary)

X_train = vectorizer_comments.fit_transform(train_comments)
X1_train = vectorizer_comments.fit_transform(train_summary)
X_test = vectorizer_summary.transform(test_comments)
X1_test = vectorizer_summary.transform(test_summary)
param_grid = {'alpha': [0.5, 1.0, 1.5, 2.0]}

clf = MultinomialNB()
grid_search = GridSearchCV(clf, param_grid, cv=5)

grid_search.fit(X_train_comments, train_ratings)

print(grid_search.best_params_)

best_clf_comments = grid_search.best_estimator_

predicted_ratings_comment = best_clf_comments.predict(X_test_comments)

clf = MultinomialNB()
grid_search = GridSearchCV(clf, param_grid, cv=5)

grid_search.fit(X_train_summary, train_ratings)

print(grid_search.best_params_)

best_clf_summary = grid_search.best_estimator_

predicted_ratings_summary = best_clf_summary.predict(X_test_summary)
inaccuracy_lst = []
min_err = 5
a_min = -1
# for k in range(0, 1000):
#     a = random.random()
# for i in range(0, 101):
predicted_ratings = []
a = 0.5
for m, n in zip(predicted_ratings_summary, predicted_ratings_comment):
    predicted_ratings.append(m*a+n*(1-a))
print(predicted_ratings)

sum = 0
for i in range(0, len(predicted_ratings)):
    sum = sum + abs(predicted_ratings[i]-test_ratings[i])
inaccuracy = sum/len(predicted_ratings)
print(inaccuracy)
# inaccuracy_lst.append(inaccuracy)
from openpyxl import load_workbook
wb = load_workbook(DATA_DIR / "app666.xlsx")
ws = wb.active
for row, entry in enumerate(predicted_ratings, start=2):
    ws.cell(row=row, column=1, value=entry)
wb.save(OUTPUT_DIR / "app666.xlsx")
#     if inaccuracy < min_err:
#         min_err = inaccuracy
#         a_min = a
# print(min_err)
# print(a_min)
